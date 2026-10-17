"""PDF and Excel exports for reports and attendee lists."""
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HEADER_FILL = PatternFill("solid", fgColor="5B3DF5")


def _ascii(value):
    return str(value if value is not None else "").encode("latin-1", "replace").decode("latin-1")


def table_xlsx(title, headers, rows, totals=None):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append(headers)
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([float(v) if hasattr(v, "quantize") else v for v in row])
    if totals:
        ws.append([float(v) if hasattr(v, "quantize") else v for v in totals])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for i, header in enumerate(headers, 1):
        width = max([len(str(header))] + [len(str(r[i - 1])) for r in rows[:200]]) + 3
        ws.column_dimensions[get_column_letter(i)].width = min(width, 45)
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def table_pdf(title, subtitle, headers, rows, totals=None):
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=28, bottomMargin=24,
                            title=title)
    styles = getSampleStyleSheet()
    data = [headers] + [[_ascii(v) for v in r] for r in rows]
    if totals:
        data.append([_ascii(v) for v in totals])
    table = Table(data, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#5B3DF5")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F3FF")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    if totals:
        style += [("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E9E5FF"))]
    table.setStyle(TableStyle(style))
    doc.build([Paragraph(_ascii(title), styles["Title"]), Paragraph(_ascii(subtitle), styles["Normal"]), Spacer(1, 10), table])
    return buf.getvalue()


def attendees_xlsx(event, tickets):
    headers = ["Ticket No.", "Attendee", "Email", "Phone", "Ticket type", "Booking ID", "Booking status", "Payment", "Check-in"]
    rows = []
    for t in tickets:
        payment = t.booking.payment
        checkin = getattr(t, "checkin", None)
        rows.append([t.number, t.attendee_name, t.booking.email, t.booking.phone, t.booking_item.ticket_type.name,
                     t.booking.booking_id, t.booking.get_status_display(), payment.get_status_display() if payment else "-",
                     timezone.localtime(checkin.checked_in_at).strftime("%Y-%m-%d %H:%M") if checkin else "Not yet"])
    return table_xlsx(f"Attendees", headers, rows)
